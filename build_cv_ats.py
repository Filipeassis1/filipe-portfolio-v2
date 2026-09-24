from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).parent / "output"
OUT.mkdir(exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.75)
section.right_margin = Cm(1.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor(25, 25, 25)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.08


def para(text, *, bold=False, size=None, before=0, after=3, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return p


def heading(text):
    para(text.upper(), bold=True, size=10, before=9, after=4, keep=True)


def bullet(text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.42)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.04
    p.add_run("- " + text)
    return p


name = para("FILIPE ASSIS", bold=True, size=17, after=2, keep=True)
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
title = para("Product Designer | UX/UI Designer", bold=True, size=10.5, after=3, keep=True)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact = para(
    "Belo Horizonte, MG | (31) 99161-7305 | contato@filipeassis.com.br | filipeassis.com.br",
    size=8.8,
    after=7,
)
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

heading("Resumo profissional")
para(
    "Product Designer e UX/UI Designer com experiência em produtos digitais B2B/SaaS, "
    "landing pages e criação de design system. Background técnico em HTML, CSS e JavaScript. "
    "Atuação da definição de requisitos ao handoff, transformando necessidades de negócio "
    "em fluxos, protótipos e interfaces no Figma, em colaboração com produto e engenharia.",
    after=3,
)

heading("Experiência profissional")
para("Gabuldev | UX/UI Designer | Jun 2026 - Atual", bold=True, after=1, keep=True)
para("Software house | Produtos B2B/SaaS", size=8.8, after=2, keep=True)
bullet(
    "Traduzi requisitos de idealizadores, gerentes de produto e stakeholders em fluxos e "
    "interfaces para produtos B2B/SaaS, acompanhando decisões com a equipe de desenvolvimento."
)
bullet(
    "Estruturei um design system no Figma com Auto Layout, componentes, variantes, "
    "propriedades e design tokens para padronizar interfaces entre produtos."
)
bullet(
    "Desenhei interfaces para diferentes produtos digitais e preparei especificações para "
    "handoff, conectando decisões de design à implementação."
)

para("Nektahub | UX/UI Designer | Jan 2026 - Set 2026", bold=True, before=5, after=1, keep=True)
para("Infoprodutos | Landing pages | Funis digitais", size=8.8, after=2, keep=True)
bullet(
    "Desenhei interfaces e páginas de vendas para infoprodutos, com foco em hierarquia "
    "visual, clareza da oferta e conversão."
)
bullet(
    "Mapeei a jornada do usuário da captação à venda e organizei fluxos de navegação "
    "para orientar a criação das páginas."
)
bullet(
    "Colaborei com especialistas e equipe de produto para transformar requisitos, "
    "benchmarking e hipóteses em interfaces consistentes."
)

para("Devsi Tecnologia | Web e UX/UI Designer | Jan 2025 - Set 2025", bold=True, before=5, after=2, keep=True)
bullet(
    "Criei protótipos navegáveis para validação interna antes do desenvolvimento, "
    "antecipando ajustes de fluxo e interface."
)
bullet(
    "Implementei interfaces em WordPress e Elementor e com HTML, CSS e JavaScript, "
    "alinhando design e desenvolvimento."
)
bullet(
    "Desenvolvi landing pages responsivas para negócios locais, com foco em hierarquia "
    "da informação e comunicação da oferta."
)

heading("Formação")
para("Análise e Desenvolvimento de Sistemas | UNINTER", bold=True, after=1, keep=True)
para("Mar 2023 - Em pausa", after=2)

heading("Cursos e certificações")
para("UX/UI e Design System | UI Boost | Em andamento, 2026", after=1)
para("User Experience (UX) | FIAP | 2023", after=1)
para("UX/UI Design e Front-End | Origamid | Em andamento, 2025", after=1)

heading("Habilidades")
para(
    "Design de produto: Product Design, UX/UI, design system, prototipação, "
    "fluxos de usuário, testes de usabilidade, handoff e design responsivo.",
    after=2,
)
para("Ferramentas: Figma, FigJam.", after=2)
para("Tecnologias: HTML, CSS, SCSS, JavaScript, Vue, Git, GitHub, WordPress e Elementor.", after=2)

doc.core_properties.title = "Currículo ATS de Filipe Assis"
doc.core_properties.subject = "Product Designer e UX/UI Designer"
doc.save(OUT / "Filipe_Assis_CV_ATS.docx")
